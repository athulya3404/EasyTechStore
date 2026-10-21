import io
import csv
import json
from decimal import Decimal
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.http import JsonResponse, HttpResponse
from django.contrib import messages
from django.db.models import Q, Sum, Count

from products.models import Product, Brand, ProductImage
from category.models import Category
from .decorators import manager_required
from .services_importer import (
    parse_product_spreadsheet_rows,
    save_single_imported_product,
)


@manager_required
def manager_products(request):
    """Display list of all products with search and filtering"""
    products = Product.objects.select_related('category', 'brand').all()
    categories = Category.objects.all()
    brands = Brand.objects.all()
    
    query = request.GET.get('q', '').strip()
    if query:
        products = products.filter(
            Q(name__icontains=query) |
            Q(description__icontains=query) |
            Q(specifications__icontains=query) |
            Q(category__name__icontains=query) |
            Q(brand__name__icontains=query)
        )
    
    category_id = request.GET.get('category')
    if category_id:
        products = products.filter(category_id=category_id)
        
    brand_id = request.GET.get('brand')
    if brand_id:
        products = products.filter(brand_id=brand_id)

    availability = request.GET.get('availability')
    if availability == 'active':
        products = products.filter(is_active=True)
    elif availability == 'inactive':
        products = products.filter(is_active=False)
    elif availability == 'in_stock':
        products = products.filter(stock__gt=10)
    elif availability == 'low_stock':
        products = products.filter(stock__gt=0, stock__lte=10)
    elif availability == 'out_of_stock':
        products = products.filter(stock=0)
    
    context = {
        'products': products,
        'categories': categories,
        'brands': brands,
        'query': query,
        'selected_category': category_id,
        'selected_brand': brand_id,
        'selected_availability': availability,
        'total_count': products.count(),
    }
    return render(request, 'manager/products.html', context)


@manager_required
def product_add(request):
    """Create a new product with all required fields"""
    categories = Category.objects.all()
    brands = Brand.objects.all()
    
    if request.method == 'POST':
        try:
            name = request.POST.get('name', '').strip()
            category_id = request.POST.get('category')
            brand_id = request.POST.get('brand') or None
            price = request.POST.get('price')
            discount = request.POST.get('discount', 0)
            stock = request.POST.get('stock', 0)
            description = request.POST.get('description', '').strip()
            specifications = request.POST.get('specifications', '').strip()
            is_active = request.POST.get('is_active') == 'on'
            image = request.FILES.get('image')
            
            if not all([name, category_id, price, stock]):
                messages.error(request, "Please fill in all required fields (Name, Category, Price, Stock).")
                return render(request, 'manager/product_add.html', {
                    'categories': categories,
                    'brands': brands,
                    'form_data': request.POST,
                })
            
            category = get_object_or_404(Category, id=category_id)
            brand = get_object_or_404(Brand, id=brand_id) if brand_id else None
            
            product = Product(
                name=name,
                category=category,
                brand=brand,
                price=Decimal(str(price)),
                discount=Decimal(str(discount)) if discount else Decimal('0'),
                stock=int(stock),
                description=description,
                specifications=specifications,
                image=image,
                is_active=is_active
            )
            product.save()
            for gfile in request.FILES.getlist('gallery_images'):
                ProductImage.objects.create(
                    product=product,
                    image=gfile,
                    alt_text=f"{product.name} angle photo"
                )
            
            messages.success(request, f"Product '{name}' created successfully!")
            return redirect('manager:products')
            
        except Exception as e:
            messages.error(request, f"Error creating product: {str(e)}")
            return render(request, 'manager/product_add.html', {
                'categories': categories,
                'brands': brands,
                'form_data': request.POST,
            })
    
    context = {
        'categories': categories,
        'brands': brands,
    }
    return render(request, 'manager/product_add.html', context)


@manager_required
def product_edit(request, pk):
    """Edit an existing product"""
    product = get_object_or_404(Product.objects.select_related('category', 'brand'), pk=pk)
    categories = Category.objects.all()
    brands = Brand.objects.all()
    
    if request.method == 'POST':
        try:
            name = request.POST.get('name', '').strip()
            category_id = request.POST.get('category')
            brand_id = request.POST.get('brand') or None
            price = request.POST.get('price')
            discount = request.POST.get('discount', 0)
            stock = request.POST.get('stock', 0)
            description = request.POST.get('description', '').strip()
            specifications = request.POST.get('specifications', '').strip()
            is_active = request.POST.get('is_active') == 'on'
            image = request.FILES.get('image')
            
            if not all([name, category_id, price, stock]):
                messages.error(request, "Please fill in all required fields (Name, Category, Price, Stock).")
                return render(request, 'manager/product_edit.html', {
                    'product': product,
                    'categories': categories,
                    'brands': brands,
                })
            
            product.name = name
            product.category = get_object_or_404(Category, id=category_id)
            product.brand = get_object_or_404(Brand, id=brand_id) if brand_id else None
            product.price = Decimal(str(price))
            product.discount = Decimal(str(discount)) if discount else Decimal('0')
            product.stock = int(stock)
            product.description = description
            product.specifications = specifications
            product.is_active = is_active
            
            if image:
                product.image = image
                
            product.save()
            for gfile in request.FILES.getlist('gallery_images'):
                ProductImage.objects.create(
                    product=product,
                    image=gfile,
                    alt_text=f"{product.name} angle photo"
                )
            messages.success(request, f"Product '{name}' updated successfully!")
            return redirect('manager:products')
            
        except Exception as e:
            messages.error(request, f"Error updating product: {str(e)}")
            return render(request, 'manager/product_edit.html', {
                'product': product,
                'categories': categories,
                'brands': brands,
            })
    
    specs_display = ""
    if isinstance(product.specifications, dict):
        specs_display = "\n".join(f"{k}: {v}" for k, v in product.specifications.items())
    elif isinstance(product.specifications, str):
        specs_display = product.specifications

    context = {
        'product': product,
        'gallery_images': product.gallery_images.all(),
        'specs_display': specs_display,
        'categories': categories,
        'brands': brands,
    }
    return render(request, 'manager/product_edit.html', context)


@manager_required
def product_detail(request, pk):
    """View product details from manager perspective"""
    product = get_object_or_404(
        Product.objects.select_related('category', 'brand').prefetch_related('reviews__user'),
        pk=pk
    )
    context = {
        'product': product,
        'reviews': product.reviews.all(),
    }
    return render(request, 'manager/product_detail.html', context)


@manager_required
def product_delete(request, pk):
    """Delete a product"""
    product = get_object_or_404(Product, pk=pk)
    
    if request.method == 'POST':
        product_name = product.name
        product.delete()
        messages.success(request, f"Product '{product_name}' deleted successfully!")
        return redirect('manager:products')
        
    return render(request, 'manager/product_confirm_delete.html', {'product': product})


@manager_required
def product_bulk_delete(request):
    """Bulk delete multiple products selected by checkboxes (supports AJAX and standard POST)"""
    if request.method == 'POST':
        is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest' or 'application/json' in request.headers.get('Accept', '')
        product_ids = request.POST.getlist('product_ids')
        if not product_ids and request.body:
            try:
                data = json.loads(request.body.decode('utf-8'))
                product_ids = data.get('product_ids', [])
            except Exception:
                product_ids = []

        if not product_ids:
            msg = "No products were selected for deletion."
            if is_ajax:
                return JsonResponse({'success': False, 'error': msg}, status=400)
            messages.warning(request, msg)
            return redirect('manager:products')

        deleted_count, _ = Product.objects.filter(id__in=product_ids).delete()
        msg = f"Successfully deleted {deleted_count} selected product(s)."
        messages.success(request, msg)

        if is_ajax:
            return JsonResponse({
                'success': True,
                'deleted_count': deleted_count,
                'message': msg,
                'redirect_url': reverse('manager:products')
            })
        return redirect('manager:products')

    return redirect('manager:products')


@manager_required
def manager_inventory(request):
    """Manage product inventory and stock levels with instant inline update"""
    products = Product.objects.select_related('category', 'brand').all()
    
    if request.method == 'POST':
        try:
            product_id = request.POST.get('product_id')
            new_stock = request.POST.get('new_stock')
            
            product = get_object_or_404(Product, id=product_id)
            old_stock = product.stock
            product.stock = max(0, int(new_stock))
            product.save()
            
            messages.success(
                request, 
                f"Stock updated for '{product.name}': {old_stock} → {product.stock} units ({product.stock_status})"
            )
        except Exception as e:
            messages.error(request, f"Error updating stock: {str(e)}")
        
        return redirect('manager:inventory')
    
    status_filter = request.GET.get('status')
    if status_filter == 'in_stock':
        products = products.filter(stock__gt=10)
    elif status_filter == 'low_stock':
        products = products.filter(stock__gt=0, stock__lte=10)
    elif status_filter == 'out_of_stock':
        products = products.filter(stock=0)
        
    query = request.GET.get('q', '').strip()
    if query:
        products = products.filter(
            Q(name__icontains=query) |
            Q(category__name__icontains=query) |
            Q(brand__name__icontains=query)
        )
    
    context = {
        'products': products,
        'selected_status': status_filter,
        'query': query,
        'total_inventory': Product.objects.aggregate(total=Sum('stock'))['total'] or 0,
        'low_stock_count': Product.objects.filter(stock__gt=0, stock__lte=10).count(),
        'out_of_stock_count': Product.objects.filter(stock=0).count(),
    }
    return render(request, 'manager/inventory.html', context)


@manager_required
def product_import(request):
    """Bulk import products with category and brand mapping from uploaded Excel (.xlsx, .xls) or CSV (.csv) file"""
    if request.method == 'POST':
        is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest' or 'application/json' in request.headers.get('Accept', '')
        uploaded_file = request.FILES.get('excel_file')
        if not uploaded_file:
            error_msg = "Please select an Excel (.xlsx) or CSV file to upload."
            if is_ajax:
                return JsonResponse({'success': False, 'error': error_msg}, status=400)
            messages.error(request, error_msg)
            return render(request, 'manager/product_import.html')

        try:
            items = parse_product_spreadsheet_rows(uploaded_file)
            created_count = 0
            already_exists_count = 0

            for item in items:
                res = save_single_imported_product(item)
                if res.get('status') == 'created':
                    created_count += 1
                elif res.get('status') == 'already_exists':
                    already_exists_count += 1

            if created_count > 0:
                msg = f"Bulk product import completed successfully! Created {created_count} new products."
                if already_exists_count > 0:
                    msg += f" {already_exists_count} existing products were preserved without being overwritten."
                messages.success(request, msg)
                if is_ajax:
                    return JsonResponse({
                        'success': True,
                        'message': msg,
                        'created_count': created_count,
                        'already_exists_count': already_exists_count,
                        'redirect_url': reverse('manager:products')
                    })
                return redirect('manager:products')
            elif already_exists_count > 0:
                msg = f"Product import completed: All {already_exists_count} products in the file already existed and were preserved without being overwritten."
                messages.info(request, msg)
                if is_ajax:
                    return JsonResponse({
                        'success': True,
                        'message': msg,
                        'created_count': 0,
                        'already_exists_count': already_exists_count,
                        'redirect_url': reverse('manager:products')
                    })
                return redirect('manager:products')
            else:
                msg = "No valid product names found in the uploaded file."
                if is_ajax:
                    return JsonResponse({'success': False, 'error': msg}, status=400)
                messages.warning(request, msg)

        except Exception as e:
            if is_ajax:
                return JsonResponse({'success': False, 'error': f"Error processing file: {str(e)}"}, status=400)
            messages.error(request, f"Error processing file: {str(e)}")

    return render(request, 'manager/product_import.html')


@manager_required
def product_import_parse(request):
    """Parses an uploaded Excel or CSV file and returns the list of product items for sequential importing"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'POST method required'}, status=405)

    uploaded_file = request.FILES.get('excel_file')
    if not uploaded_file:
        return JsonResponse({'success': False, 'error': 'Please select an Excel (.xlsx) or CSV (.csv) file.'}, status=400)

    try:
        items = parse_product_spreadsheet_rows(uploaded_file)
        if not items:
            return JsonResponse({'success': False, 'error': 'No valid product rows found in the uploaded file.'}, status=400)
        return JsonResponse({
            'success': True,
            'total_items': len(items),
            'items': items
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@manager_required
def product_import_item(request):
    """Imports an individual product item one by one to provide live sequential progress and time left"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'POST method required'}, status=405)

    try:
        if request.content_type and 'application/json' in request.content_type:
            data = json.loads(request.body.decode('utf-8'))
        else:
            data = request.POST.dict()

        result = save_single_imported_product(data)
        return JsonResponse({'success': True, 'result': result})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@manager_required
def product_sample_template(request):
    """Generates and downloads a sample file (Excel .xlsx or CSV .csv) for bulk product import including Photo URL"""
    format_type = request.GET.get('format', '').strip().lower()

    sample_rows = [
        ["Product Name", "Category", "Brand", "Price", "Discount (%)", "Stock", "Photo URL", "Specifications", "Description", "Active"],
        [
            "HP Pavilion 15",
            "Laptops",
            "HP",
            "850.00",
            "10",
            "20",
            "https://images.unsplash.com/photo-1588872657578-7efd1f1555ed?w=800",
            "Processor: Intel Core i7-1355U, RAM: 16GB DDR4, Storage: 512GB PCIe NVMe M.2 SSD, Screen: 15.6 FHD IPS",
            "Reliable all-around laptop for work, study, and daily multitasking.",
            "Yes"
        ],
        [
            "Apple MacBook Air M3",
            "Laptops",
            "Apple",
            "1099.00",
            "5",
            "25",
            "https://images.unsplash.com/photo-1517336714731-489689fd1ca8?w=800",
            "Chip: Apple M3 with 8-core CPU and 10-core GPU, RAM: 16GB Unified, Storage: 512GB SSD, Screen: 13.6 Liquid Retina",
            "Incredibly thin and fast laptop with up to 18 hours of battery life and silent fanless design.",
            "Yes"
        ],
        [
            "Samsung Galaxy S24 Ultra",
            "Smartphones",
            "Samsung",
            "1299.99",
            "8",
            "30",
            "https://images.unsplash.com/photo-1610945265064-0e34e5519bbf?w=800",
            "Display: 6.8 Dynamic AMOLED 2X 120Hz, Storage: 256GB, Camera: 200MP Quad, Battery: 5000mAh, S-Pen included",
            "Galaxy AI-powered flagship smartphone with titanium frame and built-in S Pen.",
            "Yes"
        ],
        [
            "Sony WH-1000XM5",
            "Audio & Headphones",
            "Sony",
            "399.99",
            "15",
            "40",
            "https://images.unsplash.com/photo-1505740420928-5e560c06d30e?w=800",
            "Driver: 30mm Carbon Fiber, Battery: 30 Hours with Fast Charge, NC: Auto NC Optimizer, Mics: 8 microphones",
            "Industry-leading wireless noise-cancelling headphones with crystal-clear hands-free calling.",
            "Yes"
        ],
        [
            "Logitech MX Master 3S",
            "Computer Accessories",
            "Logitech",
            "99.99",
            "0",
            "50",
            "https://images.unsplash.com/photo-1615663245857-ac93bb7c39e7?w=800",
            "Sensor: 8000 DPI Darkfield, Connectivity: Bluetooth / Logi Bolt, Battery: 70 Days Rechargeable, Quiet Clicks",
            "Performance wireless master mouse with ultra-quiet clicks and MagSpeed electromagnetic scrolling.",
            "Yes"
        ],
    ]

    if format_type == 'csv':
        output = io.StringIO()
        writer = csv.writer(output)
        for row in sample_rows:
            writer.writerow(row)
        response = HttpResponse(output.getvalue(), content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="easytech_products_template.csv"'
        return response

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"

        ws.append(sample_rows[0])
        header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        for col_num in range(1, len(sample_rows[0]) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in sample_rows[1:]:
            ws.append(row)

        ws.column_dimensions['A'].width = 28
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 14
        ws.column_dimensions['F'].width = 10
        ws.column_dimensions['G'].width = 38
        ws.column_dimensions['H'].width = 60
        ws.column_dimensions['I'].width = 50
        ws.column_dimensions['J'].width = 10

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="easytech_products_template.xlsx"'
        return response

    except (ImportError, Exception):
        output = io.StringIO()
        writer = csv.writer(output)
        for row in sample_rows:
            writer.writerow(row)
        response = HttpResponse(output.getvalue(), content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="easytech_products_template.csv"'
        return response
